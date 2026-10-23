"""Prepare the local TB chest X-ray dataset without modifying its source.

Run from the project root with:
    python src/prepare_tuberculosis_data.py

The default input is the verified sibling TB_Chest_Radiography_Database
directory. Output is written to a new sibling TB-specific directory.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import shutil
import tempfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError

CLASSES = ("Normal", "Tuberculosis")
SPLITS = ("train", "validation", "test")
SPLIT_RATIOS = {"train": 0.70, "validation": 0.15, "test": 0.15}
IMAGE_SUFFIXES = {".bmp", ".jpeg", ".jpg", ".png", ".tif", ".tiff", ".webp"}
KNOWN_DUPLICATE_PAIRS = (
    ("Tuberculosis/Tuberculosis-243.png", "Tuberculosis/Tuberculosis-244.png"),
    ("Tuberculosis/Tuberculosis-346.png", "Tuberculosis/Tuberculosis-348.png"),
    ("Tuberculosis/Tuberculosis-509.png", "Tuberculosis/Tuberculosis-510.png"),
    ("Tuberculosis/Tuberculosis-378.png", "Tuberculosis/Tuberculosis-430.png"),
    ("Tuberculosis/Tuberculosis-386.png", "Tuberculosis/Tuberculosis-405.png"),
    ("Tuberculosis/Tuberculosis-396.png", "Tuberculosis/Tuberculosis-471.png"),
)


@dataclass(frozen=True)
class ImageRecord:
    path: Path
    relative_path: str
    class_label: str
    source: str | None
    source_url: str | None
    pixel_sha256: str


def decoded_pixel_sha256(path: Path) -> str:
    """Hash canonical RGB pixels and dimensions, independent of PNG metadata."""
    with Image.open(path) as image:
        image.load()
        rgb = image.convert("RGB")
        digest = hashlib.sha256()
        digest.update(f"{rgb.width}x{rgb.height}:RGB\0".encode("ascii"))
        digest.update(rgb.tobytes())
        return digest.hexdigest()


def source_from_url(value: object) -> tuple[str | None, str | None]:
    if value is None or not str(value).strip():
        return None, None
    url = str(value).strip()
    parsed = urlparse(url)
    host = (parsed.hostname or "").casefold()
    if host.endswith("ncbi.nlm.nih.gov"):
        source = "NLM"
    elif host.endswith("kaggle.com") and "rsna-pneumonia-detection-challenge" in parsed.path:
        source = "RSNA"
    elif host.endswith("tuberculosis.by"):
        source = "Belarus portal (tuberculosis.by)"
    else:
        source = host or None
    return source, url


def normalize_filename(value: object) -> str:
    filename = Path(str(value).strip().replace("\\", "/")).name
    return filename if Path(filename).suffix else f"{filename}.png"


def read_class_metadata(dataset_root: Path, class_label: str) -> dict[str, tuple[str | None, str | None]]:
    workbook_path = dataset_root / f"{class_label}.metadata.xlsx"
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Required metadata workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValueError(f"Metadata workbook is empty: {workbook_path}")
        header_indexes = {
            str(header).strip().casefold(): index
            for index, header in enumerate(headers)
            if header is not None
        }
        filename_index = header_indexes.get("file name")
        if filename_index is None:
            raise ValueError(f"No FILE NAME column in {workbook_path}")
        url_index = header_indexes.get("url")
        metadata: dict[str, tuple[str | None, str | None]] = {}
        repeated: list[str] = []
        for row in rows:
            if filename_index >= len(row) or row[filename_index] is None:
                continue
            filename = normalize_filename(row[filename_index]).casefold()
            url = row[url_index] if url_index is not None and url_index < len(row) else None
            if filename in metadata:
                repeated.append(filename)
            metadata[filename] = source_from_url(url)
        if repeated:
            raise ValueError(f"Duplicate metadata filenames in {workbook_path}: {repeated[:10]}")
        return metadata
    finally:
        workbook.close()


def resolve_dataset_root(project_root: Path, requested: Path | None) -> Path:
    if requested is None:
        dataset_root = (project_root.parent / "TB_Chest_Radiography_Database").resolve()
    else:
        dataset_root = requested.expanduser().resolve()
    if not all((dataset_root / class_label).is_dir() for class_label in CLASSES):
        raise FileNotFoundError(
            f"Could not resolve the TB dataset at {dataset_root}; pass its actual "
            "location with --dataset-root."
        )
    for class_label in CLASSES:
        if not (dataset_root / f"{class_label}.metadata.xlsx").is_file():
            raise FileNotFoundError(f"Expected {class_label}.metadata.xlsx in {dataset_root}")
    return dataset_root


def collect_images(dataset_root: Path) -> list[ImageRecord]:
    records: list[ImageRecord] = []
    for class_label in CLASSES:
        folder = dataset_root / class_label
        metadata = read_class_metadata(dataset_root, class_label)
        paths = sorted(
            path for path in folder.iterdir()
            if path.is_file() and path.suffix.casefold() in IMAGE_SUFFIXES
        )
        disk_names = {path.name.casefold(): path for path in paths}
        missing_images = sorted(set(metadata) - set(disk_names))
        missing_metadata = sorted(set(disk_names) - set(metadata))
        if missing_images or missing_metadata:
            raise ValueError(
                f"Metadata/image mismatch for {class_label}: "
                f"metadata without image={missing_images[:10]}, "
                f"images without metadata={missing_metadata[:10]}"
            )

        for path in paths:
            try:
                pixel_hash = decoded_pixel_sha256(path)
            except (OSError, ValueError, UnidentifiedImageError) as error:
                raise ValueError(f"Unreadable source image {path}: {error}") from error
            source, source_url = metadata[path.name.casefold()]
            records.append(ImageRecord(
                path=path,
                relative_path=path.relative_to(dataset_root).as_posix(),
                class_label=class_label,
                source=source,
                source_url=source_url,
                pixel_sha256=pixel_hash,
            ))
    if not records:
        raise ValueError(f"No supported images found under {dataset_root}")
    return records


def duplicate_groups(records: list[ImageRecord]) -> list[list[ImageRecord]]:
    by_hash: defaultdict[str, list[ImageRecord]] = defaultdict(list)
    for record in records:
        by_hash[record.pixel_sha256].append(record)
    return [
        sorted(group, key=lambda record: (CLASSES.index(record.class_label), record.relative_path.casefold()))
        for group in by_hash.values()
    ]


def verify_known_pairs(records: list[ImageRecord]) -> list[dict[str, object]]:
    by_path = {record.relative_path: record for record in records}
    results = []
    for first_path, second_path in KNOWN_DUPLICATE_PAIRS:
        first = by_path.get(first_path)
        second = by_path.get(second_path)
        verified = first is not None and second is not None and first.pixel_sha256 == second.pixel_sha256
        results.append({
            "first": first_path,
            "second": second_path,
            "verified_pixel_identical": verified,
            "pixel_sha256": first.pixel_sha256 if verified and first is not None else "",
        })
    failed = [pair for pair in results if not pair["verified_pixel_identical"]]
    if failed:
        raise ValueError(f"Known duplicate verification failed: {failed}")
    return results


def allocate_stratum(records: list[ImageRecord], seed: int) -> dict[str, list[ImageRecord]]:
    shuffled = sorted(records, key=lambda record: record.relative_path.casefold())
    stratum = f"{shuffled[0].class_label}:{shuffled[0].source or 'Unknown'}"
    stable_seed = int(hashlib.sha256(f"{seed}:{stratum}".encode()).hexdigest()[:16], 16)
    random.Random(stable_seed).shuffle(shuffled)
    exact_counts = {split: len(shuffled) * ratio for split, ratio in SPLIT_RATIOS.items()}
    counts = {split: int(count) for split, count in exact_counts.items()}
    remainder = len(shuffled) - sum(counts.values())
    allocation_order = sorted(
        SPLITS,
        key=lambda split: (-(exact_counts[split] - counts[split]), SPLITS.index(split)),
    )
    for split in allocation_order[:remainder]:
        counts[split] += 1

    assigned: dict[str, list[ImageRecord]] = {split: [] for split in SPLITS}
    offset = 0
    for split in SPLITS:
        assigned[split] = shuffled[offset:offset + counts[split]]
        offset += counts[split]
    return assigned


def plan_splits(representatives: list[ImageRecord], seed: int) -> dict[str, str]:
    strata: defaultdict[tuple[str, str], list[ImageRecord]] = defaultdict(list)
    for record in representatives:
        strata[(record.class_label, record.source or "Unknown")].append(record)
    assignments: dict[str, str] = {}
    for stratum in sorted(strata):
        for split, records in allocate_stratum(strata[stratum], seed).items():
            for record in records:
                assignments[record.pixel_sha256] = split
    if len(assignments) != len(representatives):
        raise ValueError("Split planner failed to assign each representative exactly once")
    return assignments


def write_csv(path: Path, fields: tuple[str, ...], rows: list[dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def prepare(dataset_root: Path, output_dir: Path, seed: int) -> dict[str, object]:
    dataset_root = dataset_root.resolve()
    output_dir = output_dir.expanduser().resolve()
    if output_dir == dataset_root or dataset_root in output_dir.parents:
        raise ValueError("Output directory must be separate from and outside the source dataset")
    if output_dir.exists():
        raise FileExistsError(f"Output exists; refusing to overwrite it: {output_dir}")
    if not output_dir.parent.is_dir():
        raise FileNotFoundError(f"Output parent directory does not exist: {output_dir.parent}")

    records = collect_images(dataset_root)
    groups = duplicate_groups(records)
    known_pair_results = verify_known_pairs(records)
    representatives = [group[0] for group in groups]
    assignment = plan_splits(representatives, seed)
    if assignment != plan_splits(representatives, seed):
        raise ValueError("Split assignment was not reproducible")

    group_by_hash = {group[0].pixel_sha256: group for group in groups}
    record_by_path = {record.relative_path: record for record in representatives}
    duplicate_rows: list[dict[str, object]] = []
    for group in groups:
        if len(group) < 2:
            continue
        retained = group[0]
        for excluded in group[1:]:
            duplicate_rows.append({
                "duplicate_group_sha256": retained.pixel_sha256,
                "excluded_image_path": excluded.relative_path,
                "excluded_class_label": excluded.class_label,
                "excluded_source": excluded.source or "",
                "retained_representative_path": retained.relative_path,
                "retained_class_label": retained.class_label,
                "retained_source": retained.source or "",
                "assigned_split": assignment[retained.pixel_sha256],
                "duplicate_group_size": len(group),
            })

    manifest_rows: list[dict[str, object]] = []
    counts = {split: Counter({label: 0 for label in CLASSES}) for split in SPLITS}
    source_counts: defaultdict[tuple[str, str, str], int] = defaultdict(int)

    with tempfile.TemporaryDirectory(prefix=f".{output_dir.name}.staging-", dir=output_dir.parent) as temp_name:
        staging = Path(temp_name)
        for split in SPLITS:
            for label in CLASSES:
                (staging / split / label).mkdir(parents=True, exist_ok=True)

        for representative in representatives:
            split = assignment[representative.pixel_sha256]
            relative_output = Path(split) / representative.class_label / representative.path.name
            shutil.copy2(representative.path, staging / relative_output)
            group = group_by_hash[representative.pixel_sha256]
            manifest_rows.append({
                "image_path": relative_output.as_posix(),
                "original_image_path": representative.relative_path,
                "class_label": representative.class_label,
                "split": split,
                "source": representative.source or "",
                "source_url": representative.source_url or "",
                "duplicate_group_sha256": representative.pixel_sha256 if len(group) > 1 else "",
                "duplicate_group_size": len(group),
                "duplicate_handling": "retained_representative" if len(group) > 1 else "unique",
            })
            counts[split][representative.class_label] += 1
            source_counts[(split, representative.class_label, representative.source or "Unknown")] += 1

        write_csv(
            staging / "split_manifest.csv",
            ("image_path", "original_image_path", "class_label", "split", "source", "source_url",
             "duplicate_group_sha256", "duplicate_group_size", "duplicate_handling"),
            manifest_rows,
        )
        write_csv(
            staging / "duplicate_audit.csv",
            ("duplicate_group_sha256", "excluded_image_path", "excluded_class_label", "excluded_source",
             "retained_representative_path", "retained_class_label", "retained_source", "assigned_split",
             "duplicate_group_size"),
            duplicate_rows,
        )

        missing_paths: list[str] = []
        unreadable_paths: list[str] = []
        changed_pixels: list[str] = []
        output_hash_splits: defaultdict[str, set[str]] = defaultdict(set)
        for row in manifest_rows:
            image_path = str(row["image_path"])
            output_path = staging / image_path
            if not output_path.is_file():
                missing_paths.append(image_path)
                continue
            try:
                output_hash = decoded_pixel_sha256(output_path)
            except (OSError, ValueError, UnidentifiedImageError):
                unreadable_paths.append(image_path)
                continue
            source_record = record_by_path[str(row["original_image_path"])]
            if output_hash != source_record.pixel_sha256:
                changed_pixels.append(image_path)
            output_hash_splits[output_hash].add(str(row["split"]))

        leakage_groups = [
            digest for digest, split_names in output_hash_splits.items()
            if len(split_names) > 1
        ]
        image_files = [
            path for path in staging.rglob("*")
            if path.is_file() and path.suffix.casefold() in IMAGE_SUFFIXES
        ]
        actual_output_counts = {
            split: {
                label: sum(
                    1 for path in (staging / split / label).iterdir()
                    if path.is_file() and path.suffix.casefold() in IMAGE_SUFFIXES
                )
                for label in CLASSES
            }
            for split in SPLITS
        }
        retained_counts = Counter(record.class_label for record in representatives)
        class_split_ratios = {
            split: {
                label: {
                    "count": counts[split][label],
                    "ratio_of_class": round(
                        counts[split][label] / retained_counts[label], 6
                    ) if retained_counts[label] else 0,
                }
                for label in CLASSES
            }
            for split in SPLITS
        }
        split_ratio_deviations = {
            split: {
                label: round(
                    abs(class_split_ratios[split][label]["ratio_of_class"] - SPLIT_RATIOS[split]),
                    6,
                )
                for label in CLASSES
            }
            for split in SPLITS
        }
        source_count_rows = [
            {"split": split, "class_label": label, "source": source, "count": count}
            for (split, label, source), count in sorted(source_counts.items())
        ]
        source_split_ratios = []
        for label, source in sorted({(label, source) for _, label, source in source_counts}):
            source_total = sum(source_counts[(split, label, source)] for split in SPLITS)
            for split in SPLITS:
                source_split_ratios.append({
                    "split": split,
                    "class_label": label,
                    "source": source,
                    "count": source_counts[(split, label, source)],
                    "ratio_of_source": round(source_counts[(split, label, source)] / source_total, 6),
                    "requested_ratio": SPLIT_RATIOS[split],
                })
        source_ratios_within_tolerance = all(
            abs(float(row["ratio_of_source"]) - SPLIT_RATIOS[str(row["split"])]) <= 0.01
            for row in source_split_ratios
        )
        duplicate_group_details = [
            {
                "pixel_sha256": group[0].pixel_sha256,
                "members": [record.relative_path for record in group],
                "retained_representative": group[0].relative_path,
            }
            for group in groups if len(group) > 1
        ]
        validation = {
            "dataset_root": str(dataset_root),
            "output_directory": str(output_dir),
            "random_seed": seed,
            "requested_split_ratios": SPLIT_RATIOS,
            "input_class_counts": dict(Counter(record.class_label for record in records)),
            "retained_class_counts": dict(retained_counts),
            "class_counts_by_split": {split: dict(counts[split]) for split in SPLITS},
            "output_file_counts_by_split_class": actual_output_counts,
            "output_class_counts_match_manifest": actual_output_counts == {
                split: dict(counts[split]) for split in SPLITS
            },
            "class_split_ratios": class_split_ratios,
            "class_split_ratio_deviations": split_ratio_deviations,
            "class_split_ratios_within_one_percentage_point": all(
                deviation <= 0.01
                for split_deviations in split_ratio_deviations.values()
                for deviation in split_deviations.values()
            ),
            "class_source_counts_by_split": source_count_rows,
            "class_source_split_ratios": source_split_ratios,
            "source_split_ratios_within_one_percentage_point": source_ratios_within_tolerance,
            "input_image_count": len(records),
            "prepared_image_count": len(manifest_rows),
            "excluded_duplicate_image_count": len(duplicate_rows),
            "exact_pixel_duplicate_group_count": len(duplicate_group_details),
            "exact_pixel_duplicate_groups": duplicate_group_details,
            "known_duplicate_pairs_verified": known_pair_results,
            "known_duplicate_pair_count": len(known_pair_results),
            "missing_output_paths": missing_paths,
            "unreadable_output_paths": unreadable_paths,
            "changed_pixel_outputs": changed_pixels,
            "duplicate_leakage_group_count": len(leakage_groups),
            "image_file_count_matches_manifest": (
                len(image_files) == len(manifest_rows)
                and len({str(row["image_path"]) for row in manifest_rows}) == len(manifest_rows)
            ),
            "split_plan_reproducible": assignment == plan_splits(representatives, seed),
            "patient_level_separation": "not claimed; patient identifiers are unavailable",
        }
        validation["validation_passed"] = all((
            not missing_paths,
            not unreadable_paths,
            not changed_pixels,
            not leakage_groups,
            validation["image_file_count_matches_manifest"],
            validation["output_class_counts_match_manifest"],
            validation["class_split_ratios_within_one_percentage_point"],
            validation["source_split_ratios_within_one_percentage_point"],
            validation["split_plan_reproducible"],
            len(known_pair_results) == len(KNOWN_DUPLICATE_PAIRS),
        ))
        if not validation["validation_passed"]:
            raise RuntimeError(f"Prepared dataset validation failed: {validation}")

        (staging / "validation_report.json").write_text(
            json.dumps(validation, indent=2), encoding="utf-8"
        )
        staging.rename(output_dir)
    return validation


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    workspace_root = project_root.parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--dataset-root", type=Path, default=None,
        help="Actual dataset location (defaults to the inspected sibling directory)",
    )
    parser.add_argument(
        "--output-dir", type=Path,
        default=workspace_root / "TB_Chest_Radiography_Database_prepared",
        help="New output directory; existing output is never overwritten",
    )
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()
    dataset_root = resolve_dataset_root(project_root, args.dataset_root)
    print(json.dumps(prepare(dataset_root, args.output_dir, args.seed), indent=2))


if __name__ == "__main__":
    main()